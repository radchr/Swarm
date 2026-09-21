# -*- coding: utf-8 -*-
"""
Cognitive Defense Pitch Generator Template
Standard python-pptx script template for generating 18-slide MilTech pitch decks.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Color Constants
COLOR_NAVY_BLUE   = RGBColor(0x1B, 0x3A, 0x6B) # #1B3A6B
COLOR_WHITE       = RGBColor(0xFF, 0xFF, 0xFF) # #FFFFFF
COLOR_DARK_TEXT   = RGBColor(0x1E, 0x29, 0x3B) # #1E293B
COLOR_MUTED       = RGBColor(0x64, 0x74, 0x8B) # #64748B
COLOR_LIGHT_BG    = RGBColor(0xF8, 0xFA, 0xFC) # #F8FAFC
COLOR_BORDER      = RGBColor(0xCB, 0xD5, 0xE1) # #CBD5E1

FONT_HEADING = "Georgia"
FONT_BODY    = "Arial"

def build_defense_deck(output_filename="Defense_Pitch_Deck.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    def new_slide():
        s = prs.slides.add_slide(blank_layout)
        fill = s.background.fill
        fill.solid()
        fill.fore_color.rgb = COLOR_WHITE
        return s

    def add_header(s, title, kicker="COGNITIVE DEFENSE · PITCH DECK", tag="SEC // C2"):
        # Kicker
        tb_k = s.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(10), Inches(0.3))
        p_k = tb_k.text_frame.paragraphs[0]
        p_k.text = kicker.upper()
        p_k.font.name = FONT_BODY
        p_k.font.size = Pt(9.5)
        p_k.font.bold = True
        p_k.font.color.rgb = COLOR_MUTED

        # Title
        tb_t = s.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(10), Inches(0.55))
        p_t = tb_t.text_frame.paragraphs[0]
        p_t.text = title
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(20)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_NAVY_BLUE

        # Underline
        line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.25), Inches(11.733), Inches(0.025))
        line.fill.solid()
        line.fill.fore_color.rgb = COLOR_NAVY_BLUE
        line.line.color.rgb = COLOR_NAVY_BLUE

        # Tag
        tb_tag = s.shapes.add_textbox(Inches(10.5), Inches(0.4), Inches(2.0), Inches(0.3))
        p_tag = tb_tag.text_frame.paragraphs[0]
        p_tag.alignment = PP_ALIGN.RIGHT
        p_tag.text = tag
        p_tag.font.name = FONT_BODY
        p_tag.font.size = Pt(10)
        p_tag.font.bold = True
        p_tag.font.color.rgb = COLOR_NAVY_BLUE

    def add_takeaway(s, text):
        banner = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.5), Inches(11.733), Inches(0.55))
        banner.fill.solid()
        banner.fill.fore_color.rgb = COLOR_NAVY_BLUE
        banner.line.color.rgb = COLOR_NAVY_BLUE
        tf = banner.text_frame
        p = tf.paragraphs[0]
        p.text = f"  TAKEAWAY: {text}"
        p.font.name = FONT_BODY
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE

    # Add slides following patterns...
    # (Refer to pitch_structure_contract.md for exact slide content builders)

    prs.save(output_filename)
    print(f"Presentation saved to: {output_filename}")

if __name__ == "__main__":
    build_defense_deck()
